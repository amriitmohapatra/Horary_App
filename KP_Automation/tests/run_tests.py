"""
run_tests.py — Validation suite. Run after any code change:

    python run_tests.py "path/to/RA Job.xlsx"

Checks, against the two fully worked manual charts:
  1. Star/sub lord of all 24 cusps and 18 planets (exact string match).
  2. All 18 cusp-number strings, including Rahu/Ketu dispositor strings.
  3. Idempotency: processing an annotated file changes nothing.
  4. Raw-export reconstruction: stripping all manual annotations and
     re-processing reproduces the manual file cell-for-cell.
  5. Generated table: 249 subdivisions; spot-check vs the scanned PDF.
"""

from __future__ import annotations

import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

from kala_chakra.kp_core import (POINT_TO_PLANET, Chart, generate_full_table,
                     parse_longitude, star_and_sub)
from kala_chakra.process import parse_sheet, process_workbook

PASS = FAIL = 0


def check(label: str, cond: bool, detail: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
    else:
        FAIL += 1
        print(f"  FAIL: {label} {detail}")


def test_charts(path: Path) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    for sn in ("Sheet1", "Sheet3"):
        ws = wb[sn]
        chart = parse_sheet(ws)
        check(f"{sn} parsed", chart is not None)
        # manual annotations already in the file are the ground truth
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or parse_longitude(v) is None:
                    continue
                manual = re.search(r"([A-Z][a-z])/([A-Z][a-z])", v)
                if not manual:
                    continue
                _, st, su = star_and_sub(parse_longitude(v))
                check(f"{sn}!{cell.coordinate} star/sub",
                      (st, su) == manual.groups(),
                      f"manual={manual.group(0)} calc={st}/{su}")
        for r in range(38, 47):
            pt = ws[f"A{r}"].value.strip()
            manual_nums = str(ws[f"D{r}"].value).replace(" ", "")
            calc = chart.numbers_string(POINT_TO_PLANET[pt])
            check(f"{sn} {pt} numbers", calc == manual_nums,
                  f"manual={manual_nums} calc={calc}")
            manual_pow = ws[f"E{r}"].value
            if manual_pow is not None:
                manual_pow = str(manual_pow).replace(" ", "").rstrip(",")
                calc_pow = chart.powers_string(
                    POINT_TO_PLANET[pt]).replace(" ", "")
                check(f"{sn} {pt} powers", calc_pow == manual_pow,
                      f"manual={manual_pow} calc={calc_pow}")


def _diff_count(f1: Path, f2: Path,
                sheets: tuple = ("Sheet1", "Sheet3")) -> int:
    a, b = openpyxl.load_workbook(f1), openpyxl.load_workbook(f2)
    return sum(1 for sn in sheets for row in a[sn].iter_rows()
               for c in row if c.value != b[sn][c.coordinate].value)


def test_roundtrips(path: Path) -> None:
    tmp = Path(tempfile.mkdtemp())
    # idempotency
    t1 = tmp / "t1.xlsx"
    shutil.copy(path, t1)
    out1 = process_workbook(t1)
    check("idempotency (0 diffs)", _diff_count(path, out1) == 0)
    # raw-export reconstruction
    raw = openpyxl.load_workbook(path)
    strip = re.compile(r"(\d{1,2}°[A-Za-z]{2}\d{1,2}'\d{1,2}'')\s*\S+.*$")
    for sn in ("Sheet1", "Sheet3"):
        ws = raw[sn]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    m = strip.search(c.value)
                    if m:
                        c.value = c.value[:m.start(1)] + m.group(1)
        for r in range(38, 47):
            ws[f"D{r}"] = None
    t2 = tmp / "t2.xlsx"
    raw.save(t2)
    out2 = process_workbook(t2)
    check("raw-export reconstruction (0 diffs)", _diff_count(path, out2) == 0)


def test_table() -> None:
    rows = generate_full_table()
    check("249 subdivisions", len(rows) == 249, f"got {len(rows)}")
    aries = [r for r in rows if r["sign"] == "Aries"]
    check("Aries has 22 rows", len(aries) == 22)
    check("Aries row 1", (aries[0]["from"], aries[0]["to"],
                          aries[0]["star_lord"], aries[0]["sub_lord"])
          == ("00°00'00''", "00°46'40''", "Ke", "Ke"))
    check("Aries row 22", (aries[-1]["from"], aries[-1]["to"],
                           aries[-1]["star_lord"], aries[-1]["sub_lord"])
          == ("29°13'20''", "30°00'00''", "Su", "Ra"))
    ta = [r for r in rows if r["sign"] == "Taurus"][:2]
    check("Taurus handwritten row 1", (ta[0]["from"], ta[0]["to"],
                                       ta[0]["star_lord"], ta[0]["sub_lord"])
          == ("00°00'00''", "01°13'20''", "Su", "Ra"))
    check("Taurus handwritten row 2", (ta[1]["from"], ta[1]["to"],
                                       ta[1]["star_lord"], ta[1]["sub_lord"])
          == ("01°13'20''", "03°00'00''", "Su", "Ju"))


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("../RA Job.xlsx")
    test_charts(src)
    test_roundtrips(src)
    test_table()
    print(f"\n{PASS} passed, {FAIL} failed")
    sys.exit(1 if FAIL else 0)
